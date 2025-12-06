import logging
import os
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import List

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.rag.ingest import ingest_document
from app.schemas.document import DocumentItem, DocumentListResponse, DocumentUploadResponse
from app.services.financial_import import upsert_financial_statements
from database import get_db
from app.models import Document, User

router = APIRouter()
logger = logging.getLogger(__name__)

UPLOAD_DIR = Path(__file__).resolve().parent.parent / "data" / "uploads"
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
ALLOWED_EXTENSIONS = {".pdf", ".csv", ".xls", ".xlsx", ".tsv", ".txt", ".jpg", ".jpeg", ".png"}


def _ensure_upload_dir() -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def _ensure_user(db: Session, user_id: str | None) -> User | None:
    if not user_id:
        return None
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        user = User(id=user_id, nickname="ゲスト")
        db.add(user)
        db.commit()
    return user


def _extract_text(filename: str, content: bytes, mime_type: str | None) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        try:
            import pypdf

            reader = pypdf.PdfReader(BytesIO(content))
            texts = []
            for page in reader.pages[:5]:
                page_text = page.extract_text() or ""
                texts.append(page_text)
            return "\n".join(texts)
        except Exception:
            return "[PDFを受け取りました]"
    if suffix in {".csv", ".tsv"}:
        try:
            text = content.decode("utf-8", errors="ignore")
            return text[:4000]
        except Exception:
            return "[CSVを受け取りました]"
    if suffix in {".xls", ".xlsx"}:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
            sheet = wb.active
            lines: List[str] = []
            for row in list(sheet.iter_rows(values_only=True))[:20]:
                values = [str(cell) if cell is not None else "" for cell in row]
                lines.append("\t".join(values))
            return "\n".join(lines)
        except Exception:
            return "[スプレッドシートを受け取りました]"
    if mime_type and mime_type.startswith("image/"):
        return "[画像を受け取りました]"
    try:
        text = content.decode("utf-8", errors="ignore")
        return text[:4000]
    except Exception:
        return "[ファイルを受け取りました]"


# Single-file upload endpoint (chat and documents pages both rely on this).
@router.post("/documents/upload", response_model=DocumentUploadResponse)
async def upload_document(
    file: UploadFile = File(...),
    user_id: str | None = Form(None),
    company_id: str | None = Form(None),
    conversation_id: str | None = Form(None),
    doc_type: str | None = Form(None),
    period_label: str | None = Form(None),
    db: Session = Depends(get_db),
) -> DocumentUploadResponse:
    _ensure_upload_dir()
    if not (user_id or company_id or conversation_id):
        raise HTTPException(status_code=400, detail="紐づけ用のIDがありません。user_id か conversation_id を指定してください。")
    contents = await file.read()
    size_bytes = len(contents)
    if size_bytes == 0:
        raise HTTPException(status_code=400, detail="ファイルが空です。")
    if size_bytes > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="ファイルサイズは10MB以下にしてください。")

    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="サポートされていないファイル形式です。")

    saved_name = f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex}_{file.filename}"
    save_path = UPLOAD_DIR / saved_name
    with open(save_path, "wb") as f:
        f.write(contents)

    mime_type = file.content_type or "application/octet-stream"
    _ensure_user(db, user_id)
    text_content = _extract_text(file.filename or "document", contents, mime_type)
    doc = Document(
        user_id=user_id,
        company_id=company_id,
        conversation_id=conversation_id,
        filename=file.filename or "document",
        mime_type=mime_type,
        size_bytes=size_bytes,
        uploaded_at=datetime.utcnow(),
        content_text=text_content,
        doc_type=doc_type,
        period_label=period_label,
        storage_path=str(save_path),
        ingested=False,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    try:
        await ingest_document(db, doc)
    except Exception:
        logger.exception("failed to ingest document", extra={"document_id": doc.id})
        # Fail softly; document remains ingested=False

    summary = (text_content[:140] + "...") if text_content and len(text_content) > 140 else (text_content or "")
    try:
        suffix = Path(file.filename or "").suffix.lower()
        is_local_benchmark = suffix in {".xlsx", ".xlsm"} and (
            (doc_type or "").lower() in {"financial_statement", "local_benchmark"}
            or "ローカル" in (file.filename or "")
            or "benchmark" in (file.filename or "").lower()
        )
        if is_local_benchmark:
            target_company_id = company_id or "1"
            upsert_financial_statements(db, target_company_id, contents)
    except Exception:
        logger.exception("Failed to import financial statement from uploaded Excel")

    return DocumentUploadResponse(
        document_id=doc.id,
        filename=doc.filename,
        uploaded_at=doc.uploaded_at,
        summary=summary,
        storage_path=doc.storage_path,
        ingested=doc.ingested,
    )


@router.get("/documents", response_model=DocumentListResponse)
async def list_documents(user_id: str | None = None, db: Session = Depends(get_db)) -> DocumentListResponse:
    query = db.query(Document).order_by(Document.uploaded_at.desc())
    if user_id:
        query = query.filter(Document.user_id == user_id)
    docs = query.limit(50).all()
    return DocumentListResponse(
        documents=[
            DocumentItem(
                id=doc.id,
                filename=doc.filename,
                uploaded_at=doc.uploaded_at,
                size_bytes=doc.size_bytes,
                mime_type=doc.mime_type,
                content_type=doc.mime_type,
                company_id=doc.company_id,
                conversation_id=doc.conversation_id,
                doc_type=doc.doc_type,
                period_label=doc.period_label,
                storage_path=doc.storage_path,
                ingested=doc.ingested,
                content_text=doc.content_text,
            )
            for doc in docs
        ]
    )
